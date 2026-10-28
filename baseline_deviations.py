import os
import glob
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from build_change_detection_table import build_change_table

def extract_participant_num(filename):
    """Extracts integer participant number for natural sorting (par_1 -> 1, par_24 -> 24)."""
    basename = os.path.basename(filename)
    pid_str = basename.split('_aligned')[0].replace('par_', '')
    try:
        return int(pid_str)
    except ValueError:
        return 999

def build_multi_window_participant_report(
    aligned_dir=".",
    awakening_file=os.path.join(os.path.dirname(__file__), "participant_nightly_awakening", "nightly_awakenings.csv"),
    output_excel="participant_baseline_change_evaluations_multi_window.xlsx",
    output_csv=os.path.join(os.path.dirname(__file__), "change_detection_results.csv"),
):
    """
    Builds an Excel workbook with SEPARATE tabs for each participant (par_1 through par_24).
    Evaluates 8 health/sleep metrics across 4 side-by-side timeframe comparison columns:
      1. 1-Day Comparison (Prev Shift)
      2. 7-Day Baseline (1-Week)
      3. 14-Day Baseline (2-Weeks)
      4. 30-Day Baseline (1-Month)
    """
    print("Starting Multi-Window Participant Report Generation...")

    # 1. Load Awakenings Data
    awk_df = None
    if os.path.exists(awakening_file):
        awk_df = pd.read_csv(awakening_file)
        awk_df['date'] = awk_df['date'].astype(str)
        if 'participant_id' in awk_df.columns:
            awk_df['participant_id'] = awk_df['participant_id'].astype(str)

    # 2. Locate Aligned Files
    all_aligned_files = glob.glob(os.path.join(aligned_dir, "**/par_*_aligned*.csv"), recursive=True) + glob.glob(os.path.join(aligned_dir, "par_*_aligned*.csv"))
    all_aligned_files = sorted(set(all_aligned_files))

    pid_to_path = {}
    for f in all_aligned_files:
        pid = extract_participant_num(f)
        if pid != 999 and pid not in pid_to_path:
            pid_to_path[pid] = f

    if not pid_to_path:
        print("No participant aligned CSV files found.")
        return

    all_participant_rows = []

    # 3. Extract Row-Level & Nightly Summary Metrics
    for pid in sorted(pid_to_path.keys()):
        filepath = pid_to_path[pid]
        pid_str = f"par_{pid}"
        try:
            df = pd.read_csv(filepath)
        except Exception as e:
            print(f"[{pid_str}] Error reading {filepath}: {e}")
            continue

        if 'date' not in df.columns:
            continue

        df['date_str'] = df['date'].astype(str)

        for sleep_date, group in df.groupby('date_str'):
            if pd.isna(sleep_date):
                continue

            # Sleeping Heart Rate (bpm) - Excluding Awake Stage
            stage_col = 'hypnogram_class' if 'hypnogram_class' in group.columns else 'stage'
            if stage_col in group.columns and 'heart_rate' in group.columns:
                clean_stages = group[stage_col].astype(str).str.strip().str.lower()
                sleep_mask = clean_stages.isin(['light', 'deep', 'rem', 'sleep', 'n1', 'n2', 'n3']) & (group['heart_rate'] > 0)
                sleeping_hr = group.loc[sleep_mask, 'heart_rate']
                avg_hr = sleeping_hr.mean() if len(sleeping_hr) > 0 else np.nan
            else:
                avg_hr = np.nan

            # Skin Temperature Delta (°C)
            temp_delta = group['temperature_delta'].dropna().iloc[0] if 'temperature_delta' in group.columns and not group['temperature_delta'].dropna().empty else np.nan

            # Sleep Durations (converted from seconds to minutes)
            tst_min = group['night_total_sleep_seconds'].dropna().iloc[0] / 60.0 if 'night_total_sleep_seconds' in group.columns and not group['night_total_sleep_seconds'].dropna().empty else np.nan
            light_min = group['night_light_sleep_seconds'].dropna().iloc[0] / 60.0 if 'night_light_sleep_seconds' in group.columns and not group['night_light_sleep_seconds'].dropna().empty else np.nan
            deep_min = group['night_deep_sleep_seconds'].dropna().iloc[0] / 60.0 if 'night_deep_sleep_seconds' in group.columns and not group['night_deep_sleep_seconds'].dropna().empty else np.nan
            rem_min = group['night_rem_sleep_seconds'].dropna().iloc[0] / 60.0 if 'night_rem_sleep_seconds' in group.columns and not group['night_rem_sleep_seconds'].dropna().empty else np.nan
            awake_min = group['night_awake_seconds'].dropna().iloc[0] / 60.0 if 'night_awake_seconds' in group.columns and not group['night_awake_seconds'].dropna().empty else np.nan

            all_participant_rows.append({
                'participant_id': pid_str,
                'p_num': pid,
                'date': sleep_date,
                'sleeping_hr_avg_bpm': avg_hr,
                'temperature_delta_c': temp_delta,
                'total_sleep_time_min': tst_min,
                'light_sleep_time_min': light_min,
                'deep_sleep_time_min': deep_min,
                'rem_sleep_time_min': rem_min,
                'awake_time_min': awake_min
            })

    nightly_df = pd.DataFrame(all_participant_rows)

    # Merge Awakenings
    if awk_df is not None:
        nightly_df = nightly_df.merge(awk_df[['participant_id', 'date', 'awakening_count']], on=['participant_id', 'date'], how='left')
    else:
        nightly_df['awakening_count'] = np.nan

    nightly_df['date_dt'] = pd.to_datetime(nightly_df['date'])
    nightly_df = nightly_df.sort_values(['p_num', 'date_dt']).reset_index(drop=True)

    METRICS = {
        'sleeping_hr_avg_bpm': 'Sleeping Heart Rate (bpm)',
        'temperature_delta_c': 'Skin Temp Delta (°C)',
        'total_sleep_time_min': 'Total Sleep Time (min)',
        'light_sleep_time_min': 'Light Sleep Time (min)',
        'deep_sleep_time_min': 'Deep Sleep Time (min)',
        'rem_sleep_time_min': 'REM Sleep Time (min)',
        'awake_time_min': 'Awakening / Awake Time (min)',
        'awakening_count': 'Awakening Count'
    }

    TIMEFRAMES = {
        '1_day': {'days': 1, 'min_nights': 1, 'col_name': '1-Day Comparison (Prev Shift)'},
        '7_days': {'days': 7, 'min_nights': 3, 'col_name': '7-Day Baseline (1-Week)'},
        '14_days': {'days': 14, 'min_nights': 5, 'col_name': '14-Day Baseline (2-Weeks)'},
        '30_days': {'days': 30, 'min_nights': 5, 'col_name': '30-Day Baseline (1-Month)'}
    }

    # 4. Generate Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Navy Blue

    flag_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")    # Soft Red/Coral
    flag_font = Font(name="Calibri", size=10, bold=True, color="C00000")

    normal_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    alignment_center = Alignment(horizontal='center', vertical='center')
    alignment_left = Alignment(horizontal='left', vertical='center')
    alignment_right = Alignment(horizontal='right', vertical='center')

    stage_pids = sorted(nightly_df['p_num'].unique())

    for pid in stage_pids:
        pid_str = f"par_{pid}"
        p_records = nightly_df[nightly_df['p_num'] == pid].sort_values('date_dt').reset_index(drop=True)

        ws = wb.create_sheet(title=pid_str)
        ws.views.sheetView[0].showGridLines = True

        # Title Row
        ws.merge_cells('A1:G1')
        title_cell = ws['A1']
        title_cell.value = f"Participant {pid} — Personal Baseline & Anomaly Multi-Window Evaluation"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        title_cell.alignment = Alignment(horizontal='left', vertical='center')

        # Header Row
        headers = [
            "Date",
            "Metric Evaluated",
            "Current Measurement",
            "1-Day Comparison (Prev Shift)",
            "7-Day Baseline (1-Week)",
            "14-Day Baseline (2-Weeks)",
            "30-Day Baseline (1-Month)"
        ]

        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header_title
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment_center

        ws.row_dimensions[3].height = 26

        row_idx = 4

        for idx in range(len(p_records)):
            current_night = p_records.iloc[idx]
            current_date = current_night['date']

            for metric_key, metric_label in METRICS.items():
                current_val = current_night[metric_key]

                c_date_cell = ws.cell(row=row_idx, column=1, value=str(current_date))
                c_date_cell.alignment = alignment_center
                c_date_cell.font = normal_font
                c_date_cell.border = thin_border

                c_metric_cell = ws.cell(row=row_idx, column=2, value=metric_label)
                c_metric_cell.alignment = alignment_left
                c_metric_cell.font = bold_font
                c_metric_cell.border = thin_border

                c_val_str = f"{current_val:.2f}" if pd.notna(current_val) else "N/A"
                c_val_cell = ws.cell(row=row_idx, column=3, value=c_val_str)
                c_val_cell.alignment = alignment_right
                c_val_cell.font = bold_font
                c_val_cell.border = thin_border

                # Populate the 4 comparison timeframe columns
                for col_num, (tf_key, tf_config) in enumerate(TIMEFRAMES.items(), 4):
                    lookback_days = tf_config['days']
                    min_req_nights = tf_config['min_nights']

                    prior_window = p_records.iloc[max(0, idx - lookback_days):idx]
                    valid_prior = prior_window[metric_key].dropna()
                    contributing_cnt = len(valid_prior)

                    if contributing_cnt < min_req_nights:
                        cell_text = f"Insufficient Data ({contributing_cnt}/{min_req_nights} nights)"
                        is_flagged = False
                    elif pd.isna(current_val):
                        cell_text = "Missing Measurement"
                        is_flagged = False
                    else:
                        baseline_mean = valid_prior.mean()
                        baseline_std = valid_prior.std()

                        if tf_key == '1_day':
                            diff = current_val - baseline_mean
                            if abs(diff) >= 10.0 or (metric_key == 'sleeping_hr_avg_bpm' and abs(diff) >= 5.0) or (metric_key == 'temperature_delta_c' and abs(diff) >= 0.5):
                                cell_text = f"Prev: {baseline_mean:.2f} | Delta: {diff:+.2f} [FLAGGED]"
                                is_flagged = True
                            else:
                                cell_text = f"Prev: {baseline_mean:.2f} | Delta: {diff:+.2f} [Normal]"
                                is_flagged = False
                        else:
                            if pd.notna(baseline_std) and baseline_std > 0:
                                z_score = (current_val - baseline_mean) / baseline_std
                            else:
                                z_score = 0.0

                            if abs(z_score) >= 2.0:
                                cell_text = f"Base: {baseline_mean:.2f}±{baseline_std:.2f} | Z: {z_score:+.2f} [FLAGGED]"
                                is_flagged = True
                            else:
                                cell_text = f"Base: {baseline_mean:.2f}±{baseline_std:.2f} | Z: {z_score:+.2f} [Normal]"
                                is_flagged = False

                    tf_cell = ws.cell(row=row_idx, column=col_num, value=cell_text)
                    tf_cell.alignment = alignment_center
                    tf_cell.border = thin_border

                    if is_flagged:
                        tf_cell.fill = flag_fill
                        tf_cell.font = flag_font
                    else:
                        tf_cell.font = normal_font

                row_idx += 1

        # Auto-fit Column Widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value and cell.row > 1:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output_dir = os.path.dirname(output_excel)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    wb.save(output_excel)
    print(f"Successfully created workbook '{output_excel}' with {len(stage_pids)} participant tabs!")

    # Export the same baseline/change rules in a normalized format for the
    # Streamlit dashboard and downstream analysis.
    change_table = build_change_table(aligned_dir, awakening_file)
    change_table.to_csv(output_csv, index=False)
    print(f"Successfully created dashboard table '{output_csv}' with {len(change_table)} comparisons!")

if __name__ == "__main__":
    build_multi_window_participant_report()
